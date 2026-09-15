import openpyxl, json, re

wb = openpyxl.load_workbook('FTFL_Dynasty_League.xlsx', data_only=True)

TEAM_SHEETS = ['Boulder Bandits','Broad Ripple Big Horns','Denver Diamondbacks','Elkhart Express',
               'Kansas City Kaiju','Olde Town Osos','South Bend Silver Hawks','Strasbourg Soldiers',
               'Summit County Ski Bums','Wakarusa Wizards']

POSITIONS = {'QB','RB','WR','TE'}

def col_letter(idx):
    return openpyxl.utils.get_column_letter(idx)

def find_label(ws, label, max_row=45, max_col=51):
    for r in range(1, max_row+1):
        for c in range(1, max_col+1):
            v = ws.cell(r,c).value
            if isinstance(v, str) and v.strip().upper() == label.upper():
                return r,c
    return None

def parse_team(sheet_name):
    ws = wb[sheet_name]
    # 1. find position header rows in column A
    pos_headers = []
    for r in range(1, 45):
        v = ws.cell(r,1).value
        if isinstance(v,str) and v.strip().upper() in POSITIONS:
            pos_headers.append((r, v.strip().upper()))
    pos_headers.sort()

    # verify year header row (C..N) on first pos header
    year_row = pos_headers[0][0] if pos_headers else None
    year_map = {}  # col_idx -> year
    def as_year(v):
        if isinstance(v, (int, float)):
            return int(v)
        if isinstance(v, str) and re.fullmatch(r'\d{4}', v.strip()):
            return int(v.strip())
        return None

    if year_row:
        for c in range(3, 15):  # C=3 .. N=14
            yr = as_year(ws.cell(year_row, c).value)
            if yr:
                year_map[c] = yr

    players = []
    for i,(r,pos) in enumerate(pos_headers):
        next_r = pos_headers[i+1][0] if i+1 < len(pos_headers) else r+15
        for pr in range(r+1, next_r):
            name = ws.cell(pr,2).value
            if not name or not isinstance(name,str) or not name.strip():
                continue
            salaries = {}
            for c, yr in year_map.items():
                val = ws.cell(pr,c).value
                if isinstance(val,(int,float)):
                    salaries[yr] = val
            if salaries:
                players.append({'name': name.strip(), 'position': pos, 'salaries': salaries})

    def parse_side_block(label):
        entries = {}  # name -> {year: salary}
        loc = find_label(ws, label)
        if not loc:
            return entries
        r0, c0 = loc
        year_cols = {}
        # stop at the first gap — these side tables are only a handful of
        # columns wide, and scanning past them picks up unrelated tables
        # (e.g. the draft-history block) that happen to contain year-like numbers
        c = c0 + 1
        while c < c0 + 20:
            yr = as_year(ws.cell(r0, c).value)
            if not yr:
                break
            year_cols[c] = yr
            c += 1
        pr = r0+1
        blanks = 0
        while blanks < 3 and pr < r0+20:
            name = ws.cell(pr, c0).value
            if name and isinstance(name,str) and name.strip():
                sched = {}
                for c, yr in year_cols.items():
                    v = ws.cell(pr, c).value
                    if isinstance(v, (int,float)):
                        sched[yr] = v
                entries[name.strip()] = sched
                blanks = 0
            else:
                blanks += 1
            pr += 1
        return entries

    taxi_block = parse_side_block('TAXI SQUAD')
    ir_block = parse_side_block('INJURED RESERVE')

    taxi = {name: sorted(sched.keys()) for name, sched in taxi_block.items()}
    ir = {name: sorted(sched.keys()) for name, sched in ir_block.items()}

    # Players who exist ONLY in the taxi/IR side tables (not the main
    # position grid) still need a contract entry, using that side table's
    # own salary schedule. Position is unknown from this part of the sheet.
    known_names = {p['name'] for p in players}
    extra_players = []
    for name, sched in {**taxi_block, **ir_block}.items():
        if name in known_names or not sched:
            continue
        extra_players.append({
            'name': name,
            'position': '',
            'salaries': sched,
        })
        known_names.add(name)
    players.extend(extra_players)

    return {
        'sheet': sheet_name,
        'year_header_row': year_row,
        'year_map': year_map,
        'players': players,
        'taxi': taxi,
        'ir': ir,
    }

out = {}
for t in TEAM_SHEETS:
    out[t] = parse_team(t)

with open('/home/claude/parsed_contracts.json','w') as f:
    json.dump(out, f, indent=2)

for t, data in out.items():
    print(t, '| years:', sorted(data['year_map'].values()), '| players:', len(data['players']), '| taxi:', list(data['taxi'].keys()), '| ir:', list(data['ir'].keys()))
